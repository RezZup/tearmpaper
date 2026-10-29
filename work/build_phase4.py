# -*- coding: utf-8 -*-
"""Render content.py into term_paper.tex (IEEEtran + biblatex-ieee, numbered) and term_paper.docx (python-docx, OMML equations)."""
import re, json, os, sys, importlib
sys.path.insert(0, '/home/user/work')
import content as C
from bibtools import ieee_reference, ieee_names, BYKEY, write_bib, sort_key
from omml import latex_to_omml, omml_inline
from refs import R

OUT = '/home/user'
FIGDIR = '/home/user/figures'
ROMAN = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X', 'XI', 'XII', 'XIII', 'XIV', 'XV']

# ---------------------------------------------------------------- fill experiment placeholders
FILL = {}
if os.path.exists('/home/user/work/fill.json'):
    FILL = json.load(open('/home/user/work/fill.json'))

def fill(s):
    if isinstance(s, str):
        def rep(m):
            k = m.group(1)
            if k not in FILL: raise KeyError('missing fill value ' + k)
            return FILL[k]
        return re.sub(r'<<([A-Za-z0-9_]+)>>', rep, s)
    return s

def blocks():
    out = []
    for b in C.B:
        b = list(b)
        if b[0] == 'tab' and isinstance(b[4], str) and b[4].startswith('<<') and b[4].endswith('>>'): b[4] = FILL[b[4][2:-2]]
        out.append(b)
    return out

BLOCKS = blocks()

# ---------------------------------------------------------------- numbering pre-pass
NUM = {}; FIGN = TABN = EQN = 0; S1 = 0; S2 = 0
for b in BLOCKS:
    if b[0] == 'h1': S1 += 1; S2 = 0; b.append(ROMAN[S1 - 1]); NUM[b[2]] = ROMAN[S1 - 1] if b[2] else None
    elif b[0] == 'h2':
        S2 += 1; letter = 'ABCDEFGHIJ'[S2 - 1]; b.append(letter)
        if b[2]: NUM[b[2]] = f'{ROMAN[S1 - 1]}-{letter}'
    elif b[0] == 'fig': FIGN += 1; NUM[b[1]] = str(FIGN)
    elif b[0] == 'tab': TABN += 1; NUM[b[1]] = ROMAN[TABN - 1]
    elif b[0] == 'eq': EQN += 1; NUM[b[1]] = str(EQN)
NUM.pop(None, None)

TOK = re.compile(r'(\{\{.*?\}\}|\[\[.*?\]\]|⟦.*?⟧|\*\*[^*]+\*\*|\*[^*]+\*)')
CITED = []

def parse_cite(body):
    narrative = body[:2] in ('t:', 'T:'); global CAP
    CAP = body.startswith('T:')
    if narrative: body = body[2:]
    prefix = ''
    if '|' in body: prefix, body = body.split('|', 1)
    keys = [k.strip() for k in body.split(';')]
    for k in keys:
        assert k in BYKEY, 'unknown key ' + k
        CITED.append(k)
    return narrative, prefix, keys

# ---------------------------------------------------------------- LaTeX
TEX_SPECIAL = [('\\', r'\textbackslash{}'), ('&', r'\&'), ('%', r'\%'), ('$', r'\$'), ('#', r'\#'), ('_', r'\_'),
               ('{', r'\{'), ('}', r'\}'), ('~', r'\textasciitilde{}'), ('^', r'\^{}')]
TEX_UNI = {'⟨': r'$\langle$', '⟩': r'$\rangle$', '≈': r'$\approx$', '≥': r'$\geq$', 'κ': r'$\kappa$', 'α': r'$\alpha$',
           '×10⁻⁵': r'$\times 10^{-5}$', '·': r'\textperiodcentered{}', '>': r'\textgreater{}', '<': r'\textless{}', 'τ': r'$\tau$',
           '±': r'$\pm$', '†': r'$^{\dagger}$', '≤': r'$\leq$', '−': r'$-$'}

def tex_plain(s):
    for a, b in TEX_SPECIAL: s = s.replace(a, '\x00' + str(TEX_SPECIAL.index((a, b))) + '\x01')
    for i, (a, b) in enumerate(TEX_SPECIAL): s = s.replace('\x00' + str(i) + '\x01', b)
    for a, b in TEX_UNI.items(): s = s.replace(a, b)
    s = s.replace('[CLS]', '{[CLS]}').replace('[SEP]', '{[SEP]}')
    return s

def tex_inline(s):
    s = fill(s); out = []
    for part in TOK.split(s):
        if not part: continue
        if part.startswith('{{'):
            nar, pre, keys = parse_cite(part[2:-2])
            if out and out[-1].endswith(' ') and not nar and not pre: out[-1] = out[-1][:-1] + '~'
            if nar: out.append((r'\Textcite{' if CAP else r'\textcite{') + ','.join(keys) + '}')
            elif pre: out.append(pre + r',~\cite{' + ','.join(keys) + '}')
            else: out.append(r'\cite{' + ','.join(keys) + '}')
        elif part.startswith('[['):
            r = part[2:-2]; bare = r.startswith('#'); r = r.lstrip('#'); kind = r.split(':')[0]
            if bare: out.append(r'\ref{' + r + '}')
            else: out.append({'fig': 'Fig.~', 'tab': 'Table~', 'sec': 'Section~', 'eq': ''}[kind] +
                             (r'(\ref{' + r + '})' if kind == 'eq' else r'\ref{' + r + '}'))
        elif part.startswith('⟦'):
            out.append('$' + part[1:-1].split('¦')[0] + '$')
        elif part.startswith('**'): out.append(r'\textbf{' + tex_inline(part[2:-2]) + '}')
        elif part.startswith('*'): out.append(r'\emph{' + tex_inline(part[1:-1]) + '}')
        else: out.append(tex_plain(part))
    return ''.join(out)

def render_fig_tex(b):
    _, lab, path, cap, wide, width = b[:6]; env = 'figure*' if wide else 'figure'
    w = r'\textwidth' if wide else r'\columnwidth'
    frac = min(1.0, width / (7.0 if wide else 3.5))
    return (f'\\begin{{{env}}}[!t]\n\\centering\n\\includegraphics[width={frac:.2f}{w}]{{{path}.pdf}}\n'
            f'\\caption{{{tex_inline(cap)}}}\n\\label{{{lab}}}\n\\end{{{env}}}\n')

TEX_EARLY = {'fig:arch': 'sec:rw'}   # LaTeX only: emit float right after this section heading

def build_tex():
    L = []
    L.append(r'''\documentclass[conference,10pt,letterpaper]{IEEEtran}
\IEEEoverridecommandlockouts
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage{mathptmx}          % Times text and math
\usepackage[english]{babel}
\usepackage{csquotes}
\usepackage{amsmath}
\usepackage{graphicx}
\usepackage{array}
\usepackage{booktabs}
\usepackage{xurl}
\usepackage[style=ieee,backend=biber,sortcites=true,maxcitenames=2,mincitenames=1,useprefix=true]{biblatex}
\addbibresource{references.bib}
\DeclareFieldFormat{doi}{DOI\addcolon\space\ifhyperref{\href{https://doi.org/#1}{\nolinkurl{#1}}}{\nolinkurl{#1}}}
\setcounter{biburlnumpenalty}{100}\setcounter{biburlucpenalty}{100}\setcounter{biburllcpenalty}{100}
\setlength{\biburlnumskip}{0mu}\setlength{\biburlucskip}{0mu}\setlength{\biburllcskip}{0mu}
\renewcommand*{\bibfont}{\footnotesize}
\setlength{\bibhang}{1.5em}
\usepackage[hidelinks]{hyperref}
\graphicspath{{figures/}}
\begin{document}
''')
    L.append(r'\title{' + tex_plain(C.TITLE) + '}')
    a = C.AUTHOR
    L.append(r'\author{\IEEEauthorblockN{' + tex_plain(a[0]) + r'}\IEEEauthorblockA{' + r'\\{}'.join(tex_plain(x) for x in a[1:]) + '}}')
    L.append(r'\maketitle')
    L.append(r'\begin{abstract}' + tex_inline(C.ABSTRACT) + r'\end{abstract}')
    L.append(r'\begin{IEEEkeywords}' + tex_plain(C.KEYWORDS) + r'\end{IEEEkeywords}' + '\n')
    for b in BLOCKS:
        k = b[0]
        if k == 'h1':
            L.append(r'\section{' + b[1] + '}' + (r'\label{' + b[2] + '}' if b[2] else ''))
            for fl, sec in TEX_EARLY.items():
                if sec == b[2]: L.append(render_fig_tex([x for x in BLOCKS if x[0] == 'fig' and x[1] == fl][0]))
        elif k == 'h2': L.append(r'\subsection{' + b[1] + '}' + (r'\label{' + b[2] + '}' if b[2] else ''))
        elif k == 'p': L.append(tex_inline(b[1]) + '\n')
        elif k == 'list':
            L.append(r'\begin{enumerate}' + ''.join('\n\\item ' + tex_inline(x) for x in b[1]) + '\n\\end{enumerate}\n')
        elif k == 'eq':
            L.append(r'\begin{equation}' + '\n' + b[2] + '\n' + r'\label{' + b[1] + '}\n\\end{equation}')
        elif k == 'fig':
            if b[1] not in TEX_EARLY: L.append(render_fig_tex(b))
        elif k == 'tab':
            _, lab, cap, head, rows, spec, widths, wide, size = b[:9]; env = 'table*' if wide else 'table'
            body = ' & '.join(r'\textbf{' + tex_plain(h) + '}' for h in head) + r' \\' + '\n\\midrule\n'
            body += '\n'.join(' & '.join(tex_inline(c) for c in row) + r' \\' for row in rows)
            spec = re.sub(r'p\{', r'>{\\raggedright\\arraybackslash}p{', spec)
            L.append(f'\\begin{{{env}}}[!t]\n\\centering\n\\caption{{{tex_inline(cap)}}}\n\\label{{{lab}}}\n'
                     f'\\{size}\n\\setlength{{\\tabcolsep}}{{3pt}}\n\\renewcommand{{\\arraystretch}}{{1.12}}\n'
                     f'\\begin{{tabular}}{{{spec}}}\n\\toprule\n{body}\n\\bottomrule\n\\end{{tabular}}\n\\end{{{env}}}\n')
    L.append(r'\printbibliography[title={References}]')
    L.append(r'\end{document}')
    tex = '\n'.join(L)
    tex = tex.replace(r'\raggedright \textbf', r'\raggedright\textbf')
    open(f'{OUT}/term_paper.tex', 'w', encoding='utf-8').write(tex)
    write_bib(f'{OUT}/references.bib')

# ---------------------------------------------------------------- DOCX
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Times New Roman'

def set_cols(section, n, space_in=0.25):
    sp = section._sectPr
    cols = sp.find(qn('w:cols'))
    if cols is None: cols = OxmlElement('w:cols'); sp.append(cols)
    cols.set(qn('w:num'), str(n)); cols.set(qn('w:space'), str(int(space_in * 1440)))

def run(par, text, size=10, italic=False, bold=False, sub=False, sup=False, smallcaps=False):
    r = par.add_run(text); f = r.font; f.name = FONT; f.size = Pt(size); f.italic = italic; f.bold = bold
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if sub: f.subscript = True
    if sup: f.superscript = True
    if smallcaps: f.small_caps = True
    return r

def fmt(par, align='justify', first=0.0, before=0, after=0, left=0.0, hanging=0.0, keep=False):
    pf = par.paragraph_format
    par.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
    pf.first_line_indent = Inches(first) if not hanging else Inches(-hanging)
    pf.left_indent = Inches(left + hanging)
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.0
    if keep: pf.keep_with_next = True

def math_runs(par, s, size):
    """docx mini-math: _{..} subscript, ^{..} superscript; single Latin letters italic."""
    i = 0
    while i < len(s):
        if s[i] in '_^' and i + 1 < len(s) and s[i + 1] == '{':
            j = s.index('}', i); txt = s[i + 2:j]
            run(par, txt, size, italic=True, sub=s[i] == '_', sup=s[i] == '^'); i = j + 1
        else:
            j = i
            while j < len(s) and not (s[j] in '_^' and j + 1 < len(s) and s[j + 1] == '{'): j += 1
            chunk = s[i:j]
            for t in re.split(r'(\b[A-Za-z]\b)', chunk):
                if t: run(par, t, size, italic=bool(re.fullmatch(r'[A-Za-z]', t)))
            i = j

def docx_inline(par, s, size=10, base_italic=False, bold=False):
    s = fill(s)
    for part in TOK.split(s):
        if not part: continue
        if part.startswith('{{'):
            nar, pre, keys = parse_cite(part[2:-2])
            if nar:
                ps = [f"{ieee_names(k, CAP and i == 0)} [{REFNUM[k]}]" for i, k in enumerate(keys)]
                t = ps[0] if len(ps) == 1 else (' and '.join(ps) if len(ps) == 2 else ', '.join(ps[:-1]) + ', and ' + ps[-1])
            else:
                t = numlist(keys)
                if pre: t = pre + ', ' + t
            run(par, t, size, italic=base_italic, bold=bold)
        elif part.startswith('[['):
            r = part[2:-2]; bare = r.startswith('#'); r = r.lstrip('#'); kind = r.split(':')[0]; n = NUM[r]
            t = n if bare else {'fig': f'Fig. {n}', 'tab': f'Table {n}', 'sec': f'Section {n}', 'eq': f'({n})'}[kind]
            run(par, t, size, italic=base_italic, bold=bold)
        elif part.startswith('⟦'): math_runs(par, part[1:-1].split('¦')[1], size)
        elif part.startswith('**'): docx_inline(par, part[2:-2], size, base_italic, True)
        elif part.startswith('*'): docx_inline(par, part[1:-1], size, not base_italic, bold)
        else: run(par, part, size, italic=base_italic, bold=bold)

def cell_borders(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for side, val in (('top', top), ('bottom', bottom)):
        if val:
            e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(val)); e.set(qn('w:color'), '000000'); b.append(e)
    tcPr.append(b)

def cell_margins(table, lr=0.04):
    tblPr = table._tbl.tblPr; m = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), str(int(lr * 1440))); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblPr.append(m)

def bbl_order():
    """Reference order exactly as biber produced it (IEEE: order of first citation)."""
    p = f'{OUT}/term_paper.bbl'
    if os.path.exists(p):
        txt = open(p, encoding='utf-8').read()
        m = re.search(r'\\datalist\[entry\]\{none/global//[^}]*\}(.*?)\\enddatalist', txt, re.S)
        keys = re.findall(r'\\entry\{([^}]+)\}', m.group(1) if m else txt)
        assert len(keys) == len(set(keys)) == len(R), ('bbl order problem', len(keys))
        if keys: return keys
    return [r['key'] for r in sorted(R, key=sort_key)]

REFNUM = {}

def numlist(keys):
    """IEEE numeric citation as biblatex-ieee prints it: [3], [5] and ranges [3]-[5] for 3+ consecutive."""
    ns = sorted(REFNUM[k] for k in keys); out = []; i = 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1: j += 1
        if j - i >= 2: out.append(f'[{ns[i]}]\u2013[{ns[j]}]')
        else: out += [f'[{n}]' for n in ns[i:j + 1]]
        i = j + 1
    return ', '.join(out)

def add_omath(par, latex):
    from docx.oxml import parse_xml
    x = omml_inline(latex).replace('<m:oMath>', '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
                                   'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">', 1)
    par._p.append(parse_xml(x))

def build_docx():
    REFNUM.clear(); REFNUM.update({k: i + 1 for i, k in enumerate(bbl_order())})
    d = Document()
    st = d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    s0 = d.sections[0]
    s0.page_width, s0.page_height = Inches(8.5), Inches(11)
    s0.top_margin, s0.bottom_margin = Inches(0.75), Inches(1.0)
    s0.left_margin = s0.right_margin = Inches(0.625)
    set_cols(s0, 1)
    pt = d.add_paragraph(); fmt(pt, 'center', after=6); run(pt, C.TITLE, 24)
    for i, line in enumerate(C.AUTHOR):
        pa = d.add_paragraph(); fmt(pa, 'center', after=(10 if i == len(C.AUTHOR) - 1 else 0)); run(pa, line, 11 if i == 0 else 10, italic=(i in (1, 2)))
    sec = d.add_section(WD_SECTION.CONTINUOUS); set_cols(sec, 2)
    pa = d.add_paragraph(); fmt(pa, 'justify', after=4)
    run(pa, 'Abstract', 9, italic=True, bold=True); run(pa, '\u2014', 9, bold=True); docx_inline(pa, C.ABSTRACT, 9, bold=True)
    pk = d.add_paragraph(); fmt(pk, 'justify', after=6)
    run(pk, 'Index Terms', 9, italic=True, bold=True); run(pk, '\u2014', 9, bold=True); run(pk, C.KEYWORDS, 9, bold=True)
    first_after_heading = False
    for b in BLOCKS:
        k = b[0]
        if k == 'h1':
            ph = d.add_paragraph(); fmt(ph, 'center', before=8, after=4, keep=True)
            run(ph, f'{b[3]}. ', 10); run(ph, b[1].upper() if False else b[1], 10, smallcaps=True); first_after_heading = True
        elif k == 'h2':
            ph = d.add_paragraph(); fmt(ph, 'left', before=4, after=2, keep=True)
            run(ph, f'{b[3]}. {b[1]}', 10, italic=True); first_after_heading = True
        elif k == 'p':
            pp = d.add_paragraph(); fmt(pp, 'justify', first=0.14); docx_inline(pp, b[1])
        elif k == 'list':
            for i, x in enumerate(b[1], 1):
                pl = d.add_paragraph(); fmt(pl, 'justify', left=0.05, hanging=0.2); run(pl, f'{i})\t'); docx_inline(pl, x)
                pl.paragraph_format.tab_stops.add_tab_stop(Inches(0.25))
        elif k == 'eq':
            pe = d.add_paragraph(); fmt(pe, 'left', before=4, after=4)
            ts = pe.paragraph_format.tab_stops
            ts.add_tab_stop(Inches(1.75), WD_TAB_ALIGNMENT.CENTER); ts.add_tab_stop(Inches(3.5), WD_TAB_ALIGNMENT.RIGHT)
            run(pe, '\t'); add_omath(pe, b[2]); run(pe, f'\t({NUM[b[1]]})')
        elif k == 'fig':
            _, lab, path, cap, wide, width = b[:6]
            if wide: s = d.add_section(WD_SECTION.CONTINUOUS); set_cols(s, 1)
            pf = d.add_paragraph(); fmt(pf, 'center', before=4, keep=True)
            pf.add_run().add_picture(f'{FIGDIR}/{path}.png', width=Inches(width if not wide else min(width, 7.0) * 0.98))
            pc = d.add_paragraph(); fmt(pc, 'justify', after=6)
            run(pc, f'Fig. {NUM[lab]}. ', 8); docx_inline(pc, cap, 8)
            if wide: s = d.add_section(WD_SECTION.CONTINUOUS); set_cols(s, 2)
        elif k == 'tab':
            _, lab, cap, head, rows, spec, widths, wide, size = b[:9]
            fs = {'footnotesize': 7.5, 'scriptsize': 7, 'small': 8.5}[size]
            if wide: s = d.add_section(WD_SECTION.CONTINUOUS); set_cols(s, 1)
            p1 = d.add_paragraph(); fmt(p1, 'center', before=6, keep=True); run(p1, f'TABLE {NUM[lab]}', 8)
            p2 = d.add_paragraph(); fmt(p2, 'center', after=3, keep=True); docx_inline(p2, cap, 8)
            t = d.add_table(rows=1 + len(rows), cols=len(head)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False; cell_margins(t)
            tblPr = t._tbl.tblPr; lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
            tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(widths) * 1440))); tw.set(qn('w:type'), 'dxa')
            old = tblPr.find(qn('w:tblW'))
            if old is not None: tblPr.remove(old)
            tblPr.append(tw)
            for gc, wv in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(int(wv * 1440)))
            for ci, wv in enumerate(widths): t.columns[ci].width = Inches(wv)
            for ri, row in enumerate([head] + rows):
                for ci, txt in enumerate(row):
                    c = t.cell(ri, ci); c.width = Inches(widths[ci]); cp = c.paragraphs[0]
                    fmt(cp, 'left' if (spec.count('p{') == len(head) or ci == 0) else 'center')
                    if ri == 0: run(cp, txt, fs, bold=True)
                    else: docx_inline(cp, txt, fs)
                    cell_borders(c, top=(8 if ri == 0 else None), bottom=(4 if ri == 0 else (8 if ri == len(rows) else None)))
                    if ri < len(rows): cp.paragraph_format.keep_with_next = True
            sp = d.add_paragraph(); fmt(sp, after=2)
            if wide: s = d.add_section(WD_SECTION.CONTINUOUS); set_cols(s, 2)
    # references
    ph = d.add_paragraph(); fmt(ph, 'center', before=8, after=4, keep=True); run(ph, 'References', 10, smallcaps=True)
    for key in bbl_order():
        r = BYKEY[key]; pr = d.add_paragraph(); fmt(pr, 'justify', hanging=0.3, after=1)
        pr.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
        run(pr, f'[{REFNUM[key]}]\t', 8)
        for txt, it in ieee_reference(r): run(pr, txt, 8, italic=it)
    d.core_properties.title = C.TITLE; d.core_properties.author = C.AUTHOR[0]
    d.save(f'{OUT}/term_paper.docx')

if __name__ == '__main__':
    what = sys.argv[1] if len(sys.argv) > 1 else 'both'
    if what in ('tex', 'both'): build_tex()
    CITED.clear()
    if what in ('docx', 'both'): build_docx()
    cited = set(CITED); allk = {r['key'] for r in R}
    print('cited keys:', len(cited), '| uncited:', sorted(allk - cited), '| unknown:', sorted(cited - allk))
