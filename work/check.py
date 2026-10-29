"""Consistency checks for the numbered-IEEE build (tex/bib/bbl/pdf/docx)."""
import re, json, subprocess, sys
from collections import Counter
sys.path.insert(0, '/home/user/work')
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from refs import R
from bibtools import BYKEY
import content as C

tex = open('/home/user/term_paper.tex', encoding='utf-8').read()
bib = open('/home/user/references.bib', encoding='utf-8').read()
bibkeys = re.findall(r'^@\w+\{([^,]+),', bib, re.M)
texkeys = set(k.strip() for grp in re.findall(r'\\(?:[Tt]ext)?cite(?:\[[^\]]*\])*\{([^}]+)\}', tex) for k in grp.split(','))
bbl = open('/home/user/term_paper.bbl', encoding='utf-8').read()
printed = re.findall(r'\\bibitem\{([^}]+)\}', bbl)   # BibTeX/IEEEtran.bst .bbl
bibkeys = [k for k in bibkeys if k != 'IEEEexample:BSTcontrol']

d = Document('/home/user/term_paper.docx')
# body text in document order (paragraphs and tables interleaved)
seq = []
for el in d.element.body.iterchildren():
    tag = el.tag.split('}')[1]
    if tag == 'p': seq.append(Paragraph(el, d).text)
    elif tag == 'tbl':
        t = Table(el, d); seq += [c.text for row in t.rows for c in row.cells]
body = '\n'.join(seq)
i = body.rfind('\nReferences\n'); text, refs = body[:i], body[i:]

CIT = re.compile(r'\[(\d+)\](?:\u2013\[(\d+)\])?')
def cited_numbers(s):
    out = []
    for a, b in CIT.findall(s):
        a = int(a); out += list(range(a, int(b) + 1)) if b else [a]
    return out
def groups(s):  # normalised citation strings, for docx/pdf comparison
    return [re.sub(r'\s', '', g) for g in re.findall(r'\[\d+\](?:\s*(?:,|\u2013|--)\s*\[\d+\])*', s)]

docx_nums = cited_numbers(text)
first = []
for n in docx_nums:
    if n not in first: first.append(n)
dref = [p.text for p in d.paragraphs[[p.text for p in d.paragraphs].index('References') + 1:] if p.text.strip()]
ref_ok = all(t.startswith(f'[{n}]') for n, t in enumerate(dref, 1))
surname_ok = [k for n, (t, k) in enumerate(zip(dref, printed), 1)
              if BYKEY[k]['authors'][0][0].split()[-1] not in t]

pages = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', '/home/user/term_paper.pdf'], capture_output=True, text=True).stdout).group(1))
pdftxt = subprocess.run(['pdftotext', '-layout', '/home/user/term_paper.pdf', '-'], capture_output=True, text=True).stdout
pdfraw = subprocess.run(['pdftotext', '/home/user/term_paper.pdf', '-'], capture_output=True, text=True).stdout.replace('\f', '\n')
label_lines = re.findall(r'^\[(\d+)\]$', pdfraw, re.M)
_lay = pdftxt[[m.start() for m in re.finditer(r'R\s?EFERENCES', pdftxt)][-1]:]   # IEEEtran labels: start of a column line in -layout text
label_lines = sorted(set(label_lines) | set(re.findall(r'(?:^|\f|\s{3,})\[(\d+)\] +[A-Z]', _lay, re.M)), key=int)   # reference labels: pdftotext puts them on lines of their own
nolab = re.sub(r'^\[\d+\]$', '', pdfraw, flags=re.M)
j = [m.start() for m in re.finditer(r'R\s?EFERENCES', nolab)][-1]; pdf_text, pdf_refs = nolab[:j], nolab[j:]
num = {k: i + 1 for i, k in enumerate(printed)}
body_tex = tex[:tex.index('\\bibliographystyle')]
tex_nums = [num[k.strip()] for grp in re.findall(r'\\(?:[Tt]ext)?cite(?:\[[^\]]*\])*\{([^}]+)\}', body_tex) for k in grp.split(',')]
pdf_nums = cited_numbers(re.sub(r'^\[\d+\]$', '', pdfraw[:[m.start() for m in re.finditer(r'R\s?EFERENCES', pdfraw)][-1]], flags=re.M))
pdf_ref_missing = sorted(set(range(1, len(printed) + 1)) - set(int(x) for x in label_lines))
pdf_label_count = len(label_lines)
pdf_ref_surname = []   # first-author surnames must appear in bbl order through the reference list
import unicodedata
flat = unicodedata.normalize('NFC', re.sub(r'\s+', ' ', pdf_refs)); pos = 0
for n, k in enumerate(printed, 1):
    sur = unicodedata.normalize('NFC', BYKEY[k]['authors'][0][0]); q = flat.find(sur, pos)
    if q < 0: pdf_ref_surname.append((n, k))
    else: pos = q + 1
banned = [w for w in ['delve', 'moreover', 'it is important to note', "in today's", 'rapidly evolving'] if w in (text + pdfraw).lower()]
leftover = [x for x in ['{{', '}}', '[[', ']]', '<<', '>>', '\u27e6', '\u00a6', '??'] if x in body or x in pdfraw]
nfig = sum(1 for b in C.B if b[0] == 'fig'); ntab = sum(1 for b in C.B if b[0] == 'tab'); neq = sum(1 for b in C.B if b[0] == 'eq')
omath = len(d.element.body.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
peer = [r['key'] for r in R if r['type'] in ('article', 'inproceedings') and 2018 <= r['year'] <= 2026]
out = dict(pdf_pages=pages, figures=nfig, docx_images=len(d.inline_shapes), tables=ntab, docx_tables=len(d.tables),
           equations=neq, docx_omml_equations=omath,
           bib_entries=len(bibkeys), bbl_printed=len(printed), docx_reference_paragraphs=len(dref),
           tex_cited_not_in_bib=sorted(texkeys - set(bibkeys)), bib_not_cited_tex=sorted(set(bibkeys) - texkeys),
           docx_numbers_cited=sorted(set(docx_nums)) == list(range(1, len(printed) + 1)),
           docx_uncited_numbers=sorted(set(range(1, len(printed) + 1)) - set(docx_nums)),
           docx_first_citation_order_ascending=first == sorted(first),
           docx_reflist_numbered_1_to_n=ref_ok, docx_reflist_surname_mismatch=surname_ok,
           pdf_numbers_cited=sorted(set(pdf_nums)) == list(range(1, len(printed) + 1)),
           pdf_reflist_missing=pdf_ref_missing, pdf_label_lines=pdf_label_count, pdf_reflist_surname_mismatch=pdf_ref_surname,
           citation_multiset_docx_eq_tex=Counter(docx_nums) == Counter(tex_nums),
           citation_count_diff_docx_tex={n: (Counter(docx_nums)[n], Counter(tex_nums)[n]) for n in range(1, len(printed) + 1)
                                if Counter(docx_nums)[n] != Counter(tex_nums)[n]},
           banned_phrases=banned, leftover_markup=leftover,
           peer_reviewed_2018_2026=len(peer), years=sorted(r['year'] for r in R))
json.dump(out, open('/home/user/work/check.json', 'w'), indent=1)
for k, v in out.items(): print(k, ':', v)
